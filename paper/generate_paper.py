import os
import json
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

def set_cell_border(cell, **kwargs):
    """
    Set cell borders
    kwargs: top, bottom, left, right
    values: dict(sz=12, val='single', color='FF0000', space='0')
    """
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = tcPr.first_child_found_in("w:tcBorders")
    if tcBorders is None:
        tcBorders = OxmlElement('w:tcBorders')
        tcPr.append(tcBorders)
    
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        edge_data = kwargs.get(edge)
        if edge_data:
            tag = 'w:{}'.format(edge)
            element = tcBorders.find(qn(tag))
            if element is None:
                element = OxmlElement(tag)
                tcBorders.append(element)
            for key, val in edge_data.items():
                element.set(qn('w:{}'.format(key)), str(val))

def build_paper():
    # Load regenerated experimental results
    with open('results.json', 'r') as f:
        res = json.load(f)
        
    doc = docx.Document()
    
    # Configure Normal Style
    style_normal = doc.styles['Normal']
    style_normal.font.name = 'Times New Roman'
    style_normal.font.size = Pt(9.5)
    style_normal.font.color.rgb = RGBColor(0, 0, 0)
    style_normal.paragraph_format.line_spacing = 1.05
    style_normal.paragraph_format.space_after = Pt(3.5)
    style_normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    
    # Section 1: Header / Title / Authors (Full Width single column)
    sec_title = doc.sections[0]
    sec_title.top_margin = Inches(0.7)
    sec_title.bottom_margin = Inches(0.7)
    sec_title.left_margin = Inches(0.65)
    sec_title.right_margin = Inches(0.65)
    
    # Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(6)
    r_title = p_title.add_run("Coronary Artery Disease Bayesian Classification:\nDerivation, Implementation, and Empirical Validation on the UCI Cleveland Dataset")
    r_title.font.name = 'Times New Roman'
    r_title.font.size = Pt(18)
    r_title.font.bold = True
    
    # Author Block
    p_author = doc.add_paragraph()
    p_author.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_author.paragraph_format.space_after = Pt(10)
    
    r_auth = p_author.add_run("Aayush Santosh Singh\n")
    r_auth.font.name = 'Times New Roman'
    r_auth.font.size = Pt(11)
    r_auth.font.bold = True
    
    r_affil = p_author.add_run(
        "School of Computer Science and Engineering (SCOPE)\n"
        "Vellore Institute of Technology (VIT), Chennai, Tamil Nadu, India"
    )
    r_affil.font.name = 'Times New Roman'
    r_affil.font.size = Pt(9.5)
    r_affil.font.italic = True
    
    # Abstract & Keywords (Single column banner)
    p_abs = doc.add_paragraph()
    p_abs.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_abs.paragraph_format.space_after = Pt(10)
    p_abs.paragraph_format.left_indent = Inches(0.25)
    p_abs.paragraph_format.right_indent = Inches(0.25)
    
    r_abs_bold = p_abs.add_run("Abstract—")
    r_abs_bold.font.bold = True
    r_abs_bold.font.size = Pt(9)
    
    r_abs_text = p_abs.add_run(
        "Evaluating probabilistic classifiers on routine clinical indicators requires addressing heterogenous feature types, "
        "class imbalance, and statistical dependency structures. This paper presents a mathematically self-contained, leakage-free "
        "derivation and implementation of a Mixed-Feature Naive Bayes (Mixed NB) classifier for angiographic coronary artery disease (CAD) "
        "status classification. Continuous hemodynamic and physiological measurements (age, resting blood pressure, serum cholesterol, "
        "and maximum achieved heart rate) are modeled via class-conditional Gaussian likelihoods estimated through maximum likelihood parameters "
        "with an explicit variance floor. Discrete categorical indicators (chest pain type and exercise-induced angina) are modeled using "
        "categorical distributions regularized with Laplace (add-one) smoothing over bounded domain alphabets with robust zero-frequency unseen "
        "category handling. We resolve a widespread provenance and target-label encoding ambiguity in secondary distributions of the UCI Heart Disease "
        "repository, anchoring the ground truth to the canonical 303-patient Cleveland Clinic foundation cohort (164 healthy controls, 139 CAD-positive cases). "
        "The proposed classifier is evaluated against a continuous-only Gaussian Naive Bayes baseline and an L2-regularized Logistic Regression model "
        "under both an independent 80/20 stratified holdout split with 1,000 bootstrap resamples and a 10-times repeated 5-fold cross-validation protocol (50 folds). "
        "Repeated cross-validation reveals that incorporating categorical symptoms via Laplace smoothing elevates discriminative performance markedly "
        "over continuous-only modeling (+0.099 ROC-AUC gain, rising from 0.732 ± 0.052 to 0.831 ± 0.045, p < 1e-10) and matches regularized logistic regression "
        "(0.830 ± 0.049 ROC-AUC, Wilcoxon signed-rank p = 0.652). On the holdout cohort, Mixed NB achieves 80.3% accuracy [95% CI: 70.5%–90.2%], "
        "82.1% recall, and 0.891 ROC-AUC [95% CI: 0.803–0.960]. Comprehensive assumption audits via Shapiro-Wilk normality testing and within-class correlation "
        "diagnostics identify substantial violations in marginal normality (e.g., right-skewed cholesterol) and class-conditional independence (age-heart rate correlation r = -0.523), "
        "demonstrating that while classification ranking remains robust to correlation violations under zero-one loss, posterior calibration is impaired (Brier score 0.183 vs. 0.168 for logistic regression). "
        "This retrospective engineering investigation clarifies Bayesian modeling mechanics on clinical datasets without making unverified clinical validity claims."
    )
    r_abs_text.font.size = Pt(9)
    
    p_kw = doc.add_paragraph()
    p_kw.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p_kw.paragraph_format.space_after = Pt(14)
    p_kw.paragraph_format.left_indent = Inches(0.25)
    p_kw.paragraph_format.right_indent = Inches(0.25)
    
    r_kw_bold = p_kw.add_run("Index Terms—")
    r_kw_bold.font.bold = True
    r_kw_bold.font.size = Pt(9)
    r_kw_text = p_kw.add_run("Bayesian classification, Mixed Naive Bayes, coronary artery disease, Laplace smoothing, Gaussian likelihood, UCI Cleveland dataset, model calibration, diagnostic machine learning.")
    r_kw_text.font.size = Pt(9)
    
    # Section 2: Two-Column Layout for Main Text
    sec_body = doc.add_section(docx.enum.section.WD_SECTION.CONTINUOUS)
    sec_body.top_margin = Inches(0.7)
    sec_body.bottom_margin = Inches(0.7)
    sec_body.left_margin = Inches(0.65)
    sec_body.right_margin = Inches(0.65)
    
    sectPr = sec_body._sectPr
    cols = sectPr.xpath('./w:cols')
    if cols:
        cols[0].set(qn('w:num'), '2')
        cols[0].set(qn('w:space'), '720') # 0.5 inch spacing
    else:
        col = OxmlElement('w:cols')
        col.set(qn('w:num'), '2')
        col.set(qn('w:space'), '720')
        sectPr.append(col)

    def add_heading_1(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(10)
        p.paragraph_format.space_after = Pt(3)
        p.paragraph_format.keep_with_next = True
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(10)
        r.font.bold = True
        return p

    def add_heading_2(text):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        r = p.add_run(text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(9.5)
        r.font.italic = True
        return p

    def add_body_p(text):
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.05
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        r = p.add_run(text)
        r.font.name = 'Times New Roman'
        r.font.size = Pt(9.5)
        return p

    def add_fig(image_path, caption_text, width=Inches(3.3)):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(4)
        p.paragraph_format.space_after = Pt(2)
        p.paragraph_format.keep_with_next = True
        run = p.add_run()
        run.add_picture(image_path, width=width)
        
        p_cap = doc.add_paragraph()
        p_cap.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p_cap.paragraph_format.space_before = Pt(1)
        p_cap.paragraph_format.space_after = Pt(6)
        r_cap = p_cap.add_run(caption_text)
        r_cap.font.name = 'Times New Roman'
        r_cap.font.size = Pt(8)
        r_cap.font.italic = True

    # -------------------------------------------------------------
    # I. INTRODUCTION
    # -------------------------------------------------------------
    add_heading_1("I. INTRODUCTION")
    add_body_p(
        "Coronary artery disease (CAD), characterized by atherosclerotic luminal stenosis of the epicardial coronary arteries, "
        "remains the foremost cause of premature cardiovascular mortality globally. Definitive diagnosis relies on selective invasive "
        "coronary angiography, an expensive procedure carrying procedural risks. Consequently, computational decision-support tools "
        "capable of performing non-invasive diagnostic triage using routine physiological, hemodynamic, and symptomatic observations "
        "have long been pursued in biomedical informatics [1], [2]."
    )
    add_body_p(
        "In small, curated clinical registries, generative probabilistic algorithms—specifically Naive Bayes classifiers—offer compelling "
        "theoretical advantages over high-capacity black-box architectures. They estimate transparent class priors and class-conditional "
        "likelihood parameters with minimal parameter overhead, require no costly iterative hyperparameter tuning, exhibit robust behavior "
        "under small sample regimes, and yield closed-form posterior probabilities. However, typical diagnostic patient records present a mixture "
        "of continuous physical measurements (e.g., blood pressure, heart rate, cholesterol) and qualitative clinical indices (e.g., chest pain "
        "classification, presence of exercise angina). Standard off-the-shelf implementations typically force all variables into a single distribution "
        "family (e.g., treating ordinal ratings as Gaussian variables or discretizing continuous measurements into arbitrary bins), introducing "
        "severe model misspecification."
    )
    add_body_p(
        "Furthermore, academic reproducibility in machine learning benchmarks has been plagued by data provenance errors, target-label inversions, "
        "spurious duplicate removals, and train-test data leakage. In this paper, we develop, audit, and validate a Mixed-Feature Naive Bayes "
        "(Mixed NB) framework tailored to the canonical UCI Heart Disease Cleveland cohort. The primary contributions are: "
        "(1) an audited resolution of dataset provenance and ground-truth encoding, eliminating label inversion errors inherited from secondary repositories; "
        "(2) an end-to-end mathematical derivation of mixed Gaussian-categorical likelihoods with robust Laplace smoothing and unseen-category safety; "
        "(3) a strict leakage-free experimental evaluation comparing Mixed NB against continuous-only Gaussian NB and regularized Logistic Regression across "
        "an independent 80/20 holdout with bootstrap confidence intervals and a 10×5-fold repeated cross-validation scheme; and "
        "(4) empirical assumption diagnostics that elucidate the impact of normality violations and within-class correlations on classification discrimination "
        "versus probability calibration."
    )

    # -------------------------------------------------------------
    # II. RELATED WORK
    # -------------------------------------------------------------
    add_heading_1("II. RELATED WORK")
    add_body_p(
        "The application of probabilistic decision rules to cardiology dates back to the seminal work of Detrano et al. [2], who formulated "
        "a diagnostic probability algorithm across multicenter international cohorts (Cleveland, Hungary, Zurich, and Long Beach), establishing "
        "a benchmark catheterization-based threshold of >=50% diameter stenosis in at least one major vessel. Subsequent machine learning "
        "benchmarking frequently leveraged this repository [1], exploring discriminant analysis, multilayer perceptrons, decision trees, "
        "and support vector machines."
    )
    add_body_p(
        "Despite its conceptual simplicity and the known violation of conditional feature independence in living physiological systems, "
        "Naive Bayes has repeatedly demonstrated competitive empirical accuracy across medical domains. Domingos and Pazzani [3] and Rish [4] "
        "provided key theoretical explanations for this phenomenon: the zero-one misclassification loss depends strictly on the sign of the posterior "
        "log-odds ratio rather than the accuracy of the probability estimates. Thus, even when strong dependencies skew marginal joint likelihoods, "
        "the correct class boundary is preserved if the dominant class retains the maximum score. Niculescu-Mizil and Caruana [8] systematically "
        "demonstrated that while Naive Bayes preserves rank order (yielding high ROC-AUC), its output probabilities push aggressively toward 0 and 1, "
        "resulting in poor calibration and elevated Brier scores [7]. Our work builds upon these foundational insights by isolating the specific value "
        "of Laplace-smoothed categorical likelihoods when married to continuous Gaussian densities."
    )

    # -------------------------------------------------------------
    # III. DATASET AND PROBLEM FORMULATION
    # -------------------------------------------------------------
    add_heading_1("III. DATASET AND PROBLEM FORMULATION")
    add_heading_2("A. Dataset Provenance and Target Ground-Truth")
    add_body_p(
        "The Heart Disease repository at the UCI Machine Learning Repository [1] encompasses four sub-databases collected by Robert Detrano, M.D., Ph.D., "
        "and collaborators. The Cleveland Clinic Foundation database remains the standard reference. While the raw archive contains 76 attributes, "
        "published clinical studies utilize an established 14-variable subset."
    )
    add_body_p(
        "Audit of Secondary Distributions: Many internet repositories and Kaggle distributions (frequently distributed as heart.csv) contain "
        "inverted target encodings where healthy patients are marked target=1 and CAD patients target=0, or vice versa, with arbitrary column mutations. "
        "In this work, we traced data provenance directly to the primary UCI distribution (processed.cleveland.data). The authentic dataset schema dictates "
        "that the outcome variable 'num' represents angiographic CAD status: num = 0 denotes angiographic normality (<50% diameter stenosis), whereas "
        "num in {1, 2, 3, 4} denotes coronary stenosis >=50% across 1 to 4 major epicardial vessels. Binarizing via C = I(num > 0) yields exactly "
        "164 negative (healthy) records and 139 positive (CAD-present) records across N = 303 unique patient presentations."
    )
    add_body_p(
        "Spurious Duplicate Analysis: Prior draft scripts discarded one row under the assumption of an identical duplicate. Rigorous comparison across "
        "all 13 predictive attributes in the canonical archive confirms zero duplicate patient instances. Every row represents an authentic distinct clinical "
        "encounter. We therefore retain the complete cohort of N = 303 records."
    )
    add_heading_2("B. Feature Selection and Clinical Representation")
    add_body_p(
        "To rigorously examine mixed-type Bayesian mechanics without overfitting high-dimensional collinear features, we select a parsimonious six-attribute "
        "diagnostic panel comprising four continuous and two categorical features:"
    )
    add_body_p(
        "1) Age (age, years): Continuous demographic index.\n"
        "2) Resting Blood Pressure (trestbps, mm Hg): Continuous resting arterial pressure on admission.\n"
        "3) Serum Cholesterol (chol, mg/dl): Continuous fasting lipid profile.\n"
        "4) Maximum Heart Rate (thalach, beats/min): Continuous peak chronotropic response during treadmill stress testing.\n"
        "5) Chest Pain Type (cp, discrete levels 1–4): Categorical indicator (1: typical angina; 2: atypical angina; 3: non-anginal pain; 4: asymptomatic).\n"
        "6) Exercise-Induced Angina (exang, binary 0/1): Categorical indicator (0: absent; 1: provoked by exercise)."
    )
    add_body_p(
        "Diagnostic vs. Prognostic Distinction: It is vital to underscore that this formulation addresses retrospective cross-sectional disease classification "
        "(identifying existing anatomic coronary stenosis during diagnostic catheterization) rather than prospective prognostic prediction of future cardiovascular event risk."
    )

    # TABLE I: Cohort Summary
    p_t1 = doc.add_paragraph()
    p_t1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_t1.paragraph_format.space_before = Pt(6)
    p_t1.paragraph_format.space_after = Pt(2)
    p_t1.paragraph_format.keep_with_next = True
    r_t1 = p_t1.add_run("TABLE I. CLINICAL ATTRIBUTE SUMMARY AND CLASS-CONDITIONAL MEANS (N = 303)")
    r_t1.font.bold = True
    r_t1.font.size = Pt(8)

    t1 = doc.add_table(rows=7, cols=5)
    t1.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = ['Attribute', 'Clinical Meaning', 'Likelihood Model', 'Mean (C = 0)', 'Mean (C = 1)']
    for j, h in enumerate(hdr):
        cell = t1.cell(0, j)
        cell.text = h
        cell.paragraphs[0].runs[0].font.bold = True
        cell.paragraphs[0].runs[0].font.size = Pt(7.5)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_cell_border(cell, top=dict(sz=12, val='single', color='000000'), bottom=dict(sz=8, val='single', color='000000'))
    
    t1_rows = [
        ('age', 'Patient Age (years)', 'Gaussian MLE', f"{res['classmeans']['0']['age']} ± {res['class_std']['0']['age']}", f"{res['classmeans']['1']['age']} ± {res['class_std']['1']['age']}"),
        ('trestbps', 'Resting Blood Pressure (mm Hg)', 'Gaussian MLE', f"{res['classmeans']['0']['trestbps']} ± {res['class_std']['0']['trestbps']}", f"{res['classmeans']['1']['trestbps']} ± {res['class_std']['1']['trestbps']}"),
        ('chol', 'Serum Cholesterol (mg/dl)', 'Gaussian MLE', f"{res['classmeans']['0']['chol']} ± {res['class_std']['0']['chol']}", f"{res['classmeans']['1']['chol']} ± {res['class_std']['1']['chol']}"),
        ('thalach', 'Max Heart Rate Achieved (bpm)', 'Gaussian MLE', f"{res['classmeans']['0']['thalach']} ± {res['class_std']['0']['thalach']}", f"{res['classmeans']['1']['thalach']} ± {res['class_std']['1']['thalach']}"),
        ('cp', 'Chest Pain Type (4 levels)', 'Laplace Categorical', 'Levels 1–4', 'Levels 1–4'),
        ('exang', 'Exercise Induced Angina (0/1)', 'Laplace Categorical', '0: 86.0%, 1: 14.0%', '0: 45.3%, 1: 54.7%')
    ]
    for i, row in enumerate(t1_rows):
        for j, val in enumerate(row):
            cell = t1.cell(i+1, j)
            cell.text = val
            cell.paragraphs[0].runs[0].font.size = Pt(7.5)
            if j >= 3:
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_cell_border(cell, bottom=dict(sz=4, val='single', color='D3D3D3') if i < 5 else dict(sz=12, val='single', color='000000'))
    
    add_fig('figs/f2.png', "Fig. 1. Class balance (164 healthy vs. 139 CAD) and class-conditional distributions of chest pain type (cp) and exercise angina (exang). Asymptomatic pain (cp = 4) and provoked angina dominate the CAD cohort.")

    # -------------------------------------------------------------
    # IV. MATHEMATICAL FOUNDATIONS
    # -------------------------------------------------------------
    add_heading_1("IV. MATHEMATICAL FOUNDATIONS")
    add_heading_2("A. Bayesian Decision Theory and Conditional Independence")
    add_body_p(
        "Let X = (x1, x2, ..., xd)^T denote a d-dimensional feature vector, and let C in {0, 1} represent binary disease status. "
        "By Bayes' Theorem, the posterior probability of class c given observation X is expressed as:"
    )
    add_body_p(
        "P(C = c | X) = [ P(C = c) * P(X | C = c) ] / P(X) = [ P(C = c) * P(X | C = c) ] / sum_{k in {0,1}} [ P(C = k) * P(X | C = k) ]"
    )
    add_body_p(
        "Under the Naive Bayes assumption, features are assumed mutually conditionally independent given the true class c. "
        "The high-dimensional joint likelihood decomposes into the product of univariate marginal likelihoods:"
    )
    add_body_p(
        "P(X | C = c) = prod_{j=1}^d P(x_j | C = c)"
    )
    add_body_p(
        "To avoid catastrophic floating-point underflow when taking the product of numerous small probability densities, "
        "we compute the discriminant in the logarithmic domain:"
    )
    add_body_p(
        "log S(c | X) = log P(C = c) + sum_{j in NUM} log f_j(x_j | C = c) + sum_{j in CAT} log P(x_j | C = c)"
    )
    add_body_p(
        "where NUM = {age, trestbps, chol, thalach} and CAT = {cp, exang}. Binary classification decisions are determined by the log-posterior odds:"
    )
    add_body_p(
        "y_hat = 1  if  log S(1 | X) - log S(0 | X) >= tau,  else 0"
    )
    add_body_p(
        "where tau = log(theta / (1 - theta)) represents the threshold offset corresponding to operating probability threshold theta in (0, 1) (default theta = 0.5, tau = 0)."
    )
    add_heading_2("B. Continuous Features: Gaussian Likelihood and Maximum Likelihood Estimation")
    add_body_p(
        "For continuous attributes x_j in NUM, we assume class-conditional univariate Gaussian distributions: "
        "f_j(x_j | C = c) = (1 / sqrt(2 * pi * sigma_{jc}^2)) * exp{ - (x_j - mu_{jc})^2 / (2 * sigma_{jc}^2) }."
    )
    add_body_p(
        "Given N_c training records belonging to class c, the parameters mu_{jc} and sigma_{jc}^2 are derived via Maximum Likelihood Estimation (MLE):"
    )
    add_body_p(
        "mu_{jc} = (1 / N_c) * sum_{i in I_c} x_{ij},        sigma_{jc}^2 = max( (1 / N_c) * sum_{i in I_c} (x_{ij} - mu_{jc})^2,  epsilon_floor )"
    )
    add_body_p(
        "where ddof=0 corresponds strictly to the MLE variance formulation. A non-zero numerical variance floor epsilon_floor = 1e-6 is enforced "
        "to prevent division-by-zero singularities in the event of degenerate or zero-variance clinical subsets."
    )
    add_heading_2("C. Categorical Features: Multinomial Likelihood and Laplace Smoothing")
    add_body_p(
        "For qualitative discrete attributes x_j in CAT defined across finite discrete alphabet K_j, unsmoothed maximum likelihood estimates "
        "suffer from the zero-frequency problem: if category v never occurs in class c within the training split, P(x_j = v | C = c) = 0, "
        "annulling the entire multiplicative joint score. We implement additive Laplace (add-alpha) smoothing:"
    )
    add_body_p(
        "P(x_j = v | C = c) = [ N_{jc}(v) + alpha ] / [ N_c + alpha * |K_j| ]"
    )
    add_body_p(
        "where N_{jc}(v) is the empirical count of category v in class c, |K_j| is the known cardinality of valid categories, and alpha = 1.0 (Laplace's rule of succession). "
        "Crucially, to protect against out-of-vocabulary or novel categories during holdout evaluation or operational deployment, unseen levels "
        "are assigned the smoothed baseline probability alpha / (N_c + alpha * |K_j|), completely averting undefined log operations."
    )
    add_heading_2("D. Numerical Stability via Log-Sum-Exp Normalization")
    add_body_p(
        "When calibrated posterior probabilities are required for thresholding or Brier scoring, we transform the log scores back to probabilities "
        "using the stable log-sum-exp shift:"
    )
    add_body_p(
        "P(C = 1 | X) = exp(log S(1 | X) - M) / [ exp(log S(0 | X) - M) + exp(log S(1 | X) - M) ]"
    )
    add_body_p("where M = max_{c in {0,1}} log S(c | X).")

    # -------------------------------------------------------------
    # V. METHODOLOGY AND EXPERIMENTAL DESIGN
    # -------------------------------------------------------------
    add_heading_1("V. METHODOLOGY AND EXPERIMENTAL DESIGN")
    add_heading_2("A. Leakage-Free Preprocessing Pipeline")
    add_body_p(
        "Information leakage from test cohorts into training models produces inflated, ungeneralizable performance claims. To eliminate leakage: "
        "(1) All class priors P(C = c) = N_c / N_train, Gaussian parameters (mu_{jc}, sigma_{jc}^2), and categorical probability tables are estimated "
        "strictly on the training partition of each fold; (2) In logistic regression, standard scaling (Z-score normalization) and one-hot encoding columns "
        "are fitted exclusively on training records and projected onto test sets with missing dummy columns zero-filled; "
        "(3) No model selection or hyperparameter tuning was performed on test partitions; hyperparameters remained fixed at alpha = 1.0, epsilon = 1e-6."
    )
    add_heading_2("B. Validation Protocols and Uncertainty Estimation")
    add_body_p(
        "We evaluate models under two complementary paradigms:\n"
        "1) Stratified Holdout Split (80/20): N_train = 242 (131 negative, 111 positive), N_test = 61 (33 negative, 28 positive), random seed 42. "
        "To quantify sample variance on this compact test set, we construct 95% bootstrap percentile confidence intervals across 1,000 resamples.\n"
        "2) Repeated Stratified Cross-Validation (10×5-Fold): 10 independent repetitions of 5-fold stratified cross-validation (50 distinct training-testing folds), "
        "random seed 1. Because individual test folds on small samples exhibit high stochasticity, repeated cross-validation provides the primary statistical ground truth."
    )
    add_heading_2("C. Baseline Model Selection")
    add_body_p(
        "We benchmark the proposed Mixed NB against: "
        "(1) Continuous-only Gaussian Naive Bayes (Gaussian NB): Models age, trestbps, chol, thalach as Gaussians while discarding categorical symptoms. "
        "This serves as an ablation baseline to quantify the exact marginal information gain provided by Laplace-smoothed chest pain and angina indicators.\n"
        "(2) L2-Regularized Logistic Regression (LR): An established discriminative benchmark operating on standardized continuous features and one-hot encoded "
        "chest-pain and angina dummy variables (C = 1.0, max_iter = 2000)."
    )
    add_heading_2("D. Comprehensive Metric Suite")
    add_body_p(
        "Model efficacy is evaluated across seven standard clinical and machine learning metrics: Accuracy (Acc), Precision (Positive Predictive Value), "
        "Recall (Sensitivity), Specificity (True Negative Rate), F1-Score, Receiver Operating Characteristic Area Under the Curve (ROC-AUC) [6], "
        "and Brier Score Loss [7]. Differences in paired cross-validation AUC distributions between Mixed NB and Logistic Regression are analyzed using the non-parametric "
        "Wilcoxon signed-rank test."
    )

    # -------------------------------------------------------------
    # VI. RESULTS
    # -------------------------------------------------------------
    add_heading_1("VI. RESULTS")
    add_heading_2("A. Assumption Diagnostics and Distributional Audits")
    add_body_p(
        "To test the validity of the Gaussian likelihood assumption, we conducted Shapiro-Wilk tests on the continuous attributes stratified by disease status. "
        "Normality is rejected at alpha = 0.05 across multiple subsets: resting blood pressure in CAD patients (p = 0.00012), cholesterol in healthy controls (p = 2.88e-9), "
        "and maximum heart rate in healthy controls (p = 0.00054). Serum cholesterol demonstrates substantial positive skewness (skew = 1.13), with several extreme values > 400 mg/dl."
    )
    
    add_fig('figs/f1.png', "Fig. 2. Class-conditional empirical histograms and fitted Gaussian likelihood densities for continuous attributes. Maximum heart rate (thalach) achieves the cleanest class separation.")

    add_body_p(
        "Within-Class Correlation Audit: We computed the class-conditional Pearson correlation matrices to evaluate the conditional independence assumption. "
        "While most pairwise off-diagonal correlations remain modest (|r| < 0.25), a substantial negative correlation exists between age and peak heart rate "
        "among healthy patients (r = -0.523, Fig. 3). This reflects standard human chronotropic decline with age, violating strict class-conditional independence."
    )

    add_fig('figs/f3.png', "Fig. 3. Class-conditional Pearson correlation matrices. Off-diagonal values near zero support Naive Bayes; the prominent age-thalach correlation (r = -0.52) in healthy controls demonstrates conditional dependence.")

    add_heading_2("B. Predictive Performance and Ablation Analysis")
    add_body_p(
        "Table II details the empirical results across the 10×5-fold cross-validation protocol and the single 80/20 holdout split. "
        "Across all 50 folds, Mixed NB achieved a mean accuracy of 74.8% ± 4.8% and a mean ROC-AUC of 0.831 ± 0.045. "
        "Ablating the categorical features (Gaussian NB) resulted in a dramatic collapse in discriminative power: ROC-AUC dropped by 0.099 to 0.732 ± 0.052, "
        "accuracy dropped to 67.7% ± 5.0%, and sensitivity plunged from 71.4% to 60.7% (p < 1e-10). This empirical ablation proves that the Laplace-smoothed "
        "symptomatic features carry essential diagnostic signal that continuous vital signs alone cannot capture."
    )

    # TABLE II: Results Table
    p_t2 = doc.add_paragraph()
    p_t2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_t2.paragraph_format.space_before = Pt(6)
    p_t2.paragraph_format.space_after = Pt(2)
    p_t2.paragraph_format.keep_with_next = True
    r_t2 = p_t2.add_run("TABLE II. MODEL PERFORMANCE ACROSS 10x5-FOLD CV AND STRATIFIED HOLDOUT (N = 303)")
    r_t2.font.bold = True
    r_t2.font.size = Pt(8)

    t2 = doc.add_table(rows=7, cols=8)
    t2.alignment = WD_TABLE_ALIGNMENT.CENTER
    t2_hdr = ['Model', 'Protocol', 'Accuracy', 'Precision', 'Recall', 'Specificity', 'ROC-AUC', 'Brier Loss']
    for j, h in enumerate(t2_hdr):
        cell = t2.cell(0, j)
        cell.text = h
        cell.paragraphs[0].runs[0].font.bold = True
        cell.paragraphs[0].runs[0].font.size = Pt(7)
        cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        set_cell_border(cell, top=dict(sz=12, val='single', color='000000'), bottom=dict(sz=8, val='single', color='000000'))
    
    cv_m = res['cv']['Mixed NB']
    cv_g = res['cv']['Gaussian NB (numeric)']
    cv_l = res['cv']['Logistic regression']
    ho_m = res['holdout']['Mixed NB']
    ho_g = res['holdout']['Gaussian NB (numeric)']
    ho_l = res['holdout']['Logistic regression']
    
    t2_rows = [
        ('Mixed NB (Proposed)', '10x5 CV Mean±SD', f"{cv_m['acc'][0]:.3f}±{cv_m['acc'][1]:.3f}", f"{cv_m['prec'][0]:.3f}±{cv_m['prec'][1]:.3f}", f"{cv_m['rec'][0]:.3f}±{cv_m['rec'][1]:.3f}", f"{cv_m['spec'][0]:.3f}±{cv_m['spec'][1]:.3f}", f"{cv_m['auc'][0]:.3f}±{cv_m['auc'][1]:.3f}", f"{cv_m['brier'][0]:.3f}±{cv_m['brier'][1]:.3f}"),
        ('Gaussian NB (numeric)', '10x5 CV Mean±SD', f"{cv_g['acc'][0]:.3f}±{cv_g['acc'][1]:.3f}", f"{cv_g['prec'][0]:.3f}±{cv_g['prec'][1]:.3f}", f"{cv_g['rec'][0]:.3f}±{cv_g['rec'][1]:.3f}", f"{cv_g['spec'][0]:.3f}±{cv_g['spec'][1]:.3f}", f"{cv_g['auc'][0]:.3f}±{cv_g['auc'][1]:.3f}", f"{cv_g['brier'][0]:.3f}±{cv_g['brier'][1]:.3f}"),
        ('Logistic Regression', '10x5 CV Mean±SD', f"{cv_l['acc'][0]:.3f}±{cv_l['acc'][1]:.3f}", f"{cv_l['prec'][0]:.3f}±{cv_l['prec'][1]:.3f}", f"{cv_l['rec'][0]:.3f}±{cv_l['rec'][1]:.3f}", f"{cv_l['spec'][0]:.3f}±{cv_l['spec'][1]:.3f}", f"{cv_l['auc'][0]:.3f}±{cv_l['auc'][1]:.3f}", f"{cv_l['brier'][0]:.3f}±{cv_l['brier'][1]:.3f}"),
        ('Mixed NB (Proposed)', 'Holdout [95% CI]', f"{ho_m['acc']:.3f} [{res['holdout_ci']['Mixed NB']['acc'][0]:.3f}, {res['holdout_ci']['Mixed NB']['acc'][1]:.3f}]", f"{ho_m['prec']:.3f}", f"{ho_m['rec']:.3f}", f"{ho_m['spec']:.3f}", f"{ho_m['auc']:.3f} [{res['holdout_ci']['Mixed NB']['auc'][0]:.3f}, {res['holdout_ci']['Mixed NB']['auc'][1]:.3f}]", f"{ho_m['brier']:.3f}"),
        ('Gaussian NB (numeric)', 'Holdout [95% CI]', f"{ho_g['acc']:.3f} [{res['holdout_ci']['Gaussian NB (numeric)']['acc'][0]:.3f}, {res['holdout_ci']['Gaussian NB (numeric)']['acc'][1]:.3f}]", f"{ho_g['prec']:.3f}", f"{ho_g['rec']:.3f}", f"{ho_g['spec']:.3f}", f"{ho_g['auc']:.3f} [{res['holdout_ci']['Gaussian NB (numeric)']['auc'][0]:.3f}, {res['holdout_ci']['Gaussian NB (numeric)']['auc'][1]:.3f}]", f"{ho_g['brier']:.3f}"),
        ('Logistic Regression', 'Holdout [95% CI]', f"{ho_l['acc']:.3f} [{res['holdout_ci']['Logistic regression']['acc'][0]:.3f}, {res['holdout_ci']['Logistic regression']['acc'][1]:.3f}]", f"{ho_l['prec']:.3f}", f"{ho_l['rec']:.3f}", f"{ho_l['spec']:.3f}", f"{ho_l['auc']:.3f} [{res['holdout_ci']['Logistic regression']['auc'][0]:.3f}, {res['holdout_ci']['Logistic regression']['auc'][1]:.3f}]", f"{ho_l['brier']:.3f}")
    ]
    for i, row in enumerate(t2_rows):
        for j, val in enumerate(row):
            cell = t2.cell(i+1, j)
            cell.text = val
            cell.paragraphs[0].runs[0].font.size = Pt(7)
            if j >= 2:
                cell.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_cell_border(cell, bottom=dict(sz=4, val='single', color='D3D3D3') if i != 2 and i != 5 else dict(sz=12, val='single', color='000000'))

    add_body_p(
        "Comparison with Logistic Regression: Mixed NB achieved statistical parity with regularized Logistic Regression in discrimination. "
        "The paired mean fold AUC difference was +0.0011 ± 0.0188 (Wilcoxon signed-rank p = 0.652, confirming no statistically significant difference). "
        "Mixed NB achieved slightly higher recall (71.4% vs. 70.7%) and comparable specificity (77.6% vs. 78.8%)."
    )

    add_fig('figs/f5.png', "Fig. 4. Mean cross-validation performance comparison across 50 folds (error bars denote ±1 standard deviation). Mixed NB matches Logistic Regression while substantially exceeding Gaussian NB.")

    add_heading_2("C. Holdout Classification and Confusion Analysis")
    add_body_p(
        "On the independent 61-patient holdout set (28 CAD cases, 33 controls), Mixed NB correctly identified 23 true positives and 26 true negatives, "
        "yielding 80.3% accuracy [95% CI: 70.5%–90.2%], 82.1% recall, 78.8% specificity, and 0.891 ROC-AUC [95% CI: 0.803–0.960] (Fig. 5). "
        "Logistic regression obtained 82.0% accuracy, 78.6% recall, and 0.898 ROC-AUC. The overlapping bootstrap confidence intervals underscore that "
        "a single 61-record test set lacks the statistical power to resolve micro-differences between the models, affirming the necessity of repeated cross-validation."
    )

    add_fig('figs/f6.png', "Fig. 5. Confusion matrices on the stratified holdout cohort (N = 61; positive = CAD). Mixed NB achieved 23 TP, 5 FN, 7 FP, and 26 TN.")

    add_heading_2("D. Discrimination, Calibration, and Operating Thresholds")
    add_body_p(
        "Fig. 6(a) displays the pooled out-of-fold ROC curves. Mixed NB and Logistic Regression curves trace almost identical trajectories (pooled AUC 0.831 vs 0.830). "
        "However, the reliability calibration diagram (Fig. 6b) reveals a stark structural divergence. Logistic regression tracks the diagonal calibration line closely, "
        "achieving a superior Brier score of 0.168 ± 0.026. Conversely, Mixed NB exhibits an S-shaped calibration distortion, overestimating probabilities in intermediate "
        "ranges and producing a higher Brier score of 0.183 ± 0.032. This aligns precisely with Rish [4] and Niculescu-Mizil [8]: the product of non-independent likelihoods "
        "exaggerates posterior odds without harming rank ordering."
    )

    add_fig('figs/f4.png', "Fig. 6. (a) Pooled out-of-fold ROC curves (Mixed NB and Logistic Regression achieve identical 0.831/0.830 AUC); (b) Reliability calibration diagrams illustrating superior calibration for Logistic Regression.")

    add_body_p(
        "Operating Threshold Sweep: In clinical triage, false negatives (missed ischemic heart disease) carry far greater medical consequence than false positives "
        "(recommending unnecessary confirmatory testing). Fig. 7 illustrates the sensitivity-specificity tradeoff across operating thresholds theta in [0.05, 0.95]. "
        "Lowering the decision threshold to theta = 0.35 elevates sensitivity to 85.6% while maintaining 68.3% specificity, illustrating practical adaptability."
    )

    add_fig('figs/f7.png', "Fig. 7. Operating threshold sweep for Mixed NB. Lowering the threshold below the default 0.5 trades specificity to achieve elevated clinical recall.")

    # -------------------------------------------------------------
    # VII. DISCUSSION
    # -------------------------------------------------------------
    add_heading_1("VII. DISCUSSION")
    add_body_p(
        "Mechanisms of Performance Gains: The marked +0.099 AUC elevation when integrating chest-pain type and exercise-induced angina via Laplace smoothing "
        "demonstrates the power of structured generative models. Symptomatic angina reflects dynamic myocardial oxygen supply-demand mismatches that resting "
        "vital signs cannot capture. Laplace smoothing provides an optimal inductive bias in small sample regimes, preventing zero-probability traps without "
        "requiring high-variance parameter tuning."
    )
    add_body_p(
        "Why Independence Violations Do Not Degrade Classification: Living biological organisms exhibit extensive physiological coupling; resting blood pressure, "
        "cardiac chronotropic capacity, and age are intimately correlated. Despite a correlation of r = -0.523 between age and peak heart rate in controls, "
        "Mixed NB matches the discriminative ability of regularized Logistic Regression. This empirical result confirms the zero-one optimality theorem of Domingos "
        "and Pazzani [3]: conditional independence is a sufficient but strictly non-necessary condition for optimal Bayesian classification. When dependencies distribute "
        "symmetrically across classes, the classification decision boundary remains invariant."
    )
    add_body_p(
        "The Calibration Cost: While ranking is robust, the probability magnitudes are distorted. Naive Bayes implicitly counts redundant shared information multiple times, "
        "driving posterior estimates toward certainty (0 or 1). In clinical applications requiring risk stratification or decision-theoretic cost weighting, "
        "post-hoc probability calibration (e.g., Platt scaling or isotonic regression) is essential."
    )

    # -------------------------------------------------------------
    # VIII. LIMITATIONS AND ETHICAL CONSIDERATIONS
    # -------------------------------------------------------------
    add_heading_1("VIII. LIMITATIONS AND ETHICAL CONSIDERATIONS")
    add_body_p(
        "Several methodological limitations must be stated:\n"
        "1) Single-Center Historical Data: The cohort consists of 303 patients evaluated at the Cleveland Clinic Foundation in the late 1980s. Modern pharmacotherapy "
        "(e.g., statins, ACE inhibitors, novel beta-blockers) and diagnostic imaging standards have evolved dramatically. External validation across modern cohorts is necessary.\n"
        "2) Binary Binarization Loss: Collapsing the multi-vessel score (num in {1, 2, 3, 4}) into a binary indicator ignores anatomical severity (single-vessel vs. triple-vessel or left main disease).\n"
        "3) Subselected Predictor Panel: We deliberately restricted modeling to six routine indicators to examine mixed-type mechanics. Incorporating remaining predictors "
        "(ST depression, fluoroscopy vessel counts, thallium scintigraphy) would yield higher absolute discrimination but involves substantial missingness (e.g., ca and thal).\n"
        "4) Educational Scope and Non-Clinical Disclaimer: This paper represents an engineering and computational statistics study. The resulting models have not underwent "
        "prospective clinical trials and must not be utilized for autonomous clinical triage or therapeutic interventions."
    )

    # -------------------------------------------------------------
    # IX. CONCLUSION
    # -------------------------------------------------------------
    add_heading_1("IX. CONCLUSION")
    add_body_p(
        "This work formulated, implemented, and empirically audited a Mixed-Feature Naive Bayes classifier designed for coronary artery disease classification "
        "on the UCI Cleveland database. By resolving historical target-encoding ambiguities and establishing a leakage-free pipeline, we demonstrated that "
        "Laplace-smoothed categorical symptoms yield a statistically substantial improvement (+0.099 AUC) over continuous-only modeling, achieving 0.831 CV ROC-AUC "
        "and matching L2-regularized logistic regression. Distributional diagnostics confirmed that while marginal non-normality and conditional dependencies do not "
        "impair classification discrimination, they produce calibration distortion. Future work will investigate non-parametric kernel density estimation, "
        "Bayesian network architectures with learned dependence structures, and post-hoc calibration."
    )

    # -------------------------------------------------------------
    # REPRODUCIBILITY STATEMENT
    # -------------------------------------------------------------
    add_heading_1("REPRODUCIBILITY STATEMENT")
    add_body_p(
        "All code, processed datasets, figures, and numerical metrics are completely reproducible. "
        "Experiments were executed in Python 3.10 using NumPy, pandas, SciPy, scikit-learn, and Matplotlib. "
        "Random seeds: Stratified holdout split seed = 42; Repeated cross-validation seed = 1; Bootstrap resample seed = 0. "
        "Hyperparameters: Laplace smoothing alpha = 1.0; Variance floor epsilon = 1e-6; Decision threshold theta = 0.5. "
        "The complete source pipeline is archived in experiments.py and results.json."
    )

    # -------------------------------------------------------------
    # REFERENCES
    # -------------------------------------------------------------
    add_heading_1("REFERENCES")
    refs = [
        "[1] A. Janosi, W. Steinbrunn, M. Pfisterer, and R. Detrano, \"Heart Disease,\" UCI Machine Learning Repository, 1989. [Online]. Available: https://doi.org/10.24432/C52P4X.",
        "[2] R. Detrano, A. Janosi, W. Steinbrunn, M. Pfisterer, J. Schmid, S. Sandhu, K. Guppy, S. Lee, and V. Froelicher, \"International application of a new probability algorithm for the diagnosis of coronary artery disease,\" The American Journal of Cardiology, vol. 64, no. 5, pp. 304–310, 1989.",
        "[3] P. Domingos and M. Pazzani, \"On the optimality of the simple Bayesian classifier under zero-one loss,\" Machine Learning, vol. 29, no. 2, pp. 103–130, 1997.",
        "[4] I. Rish, \"An empirical study of the naive Bayes classifier,\" in Proc. IJCAI-2001 Workshop on Empirical Methods in Artificial Intelligence, Seattle, WA, USA, 2001, pp. 41–46.",
        "[5] F. Pedregosa, G. Varoquaux, A. Gramfort, V. Michel, B. Thirion, O. Grisel, M. Blondel, P. Prettenhofer, R. Weiss, V. Dubourg, J. Vanderplas, A. Passos, D. Cournapeau, M. Brucher, M. Perrot, and E. Duchesnay, \"Scikit-learn: Machine learning in Python,\" Journal of Machine Learning Research, vol. 12, pp. 2825–2830, 2011.",
        "[6] T. Fawcett, \"An introduction to ROC analysis,\" Pattern Recognition Letters, vol. 27, no. 8, pp. 861–874, 2006.",
        "[7] G. W. Brier, \"Verification of forecasts expressed in terms of probability,\" Monthly Weather Review, vol. 78, no. 1, pp. 1–3, 1950.",
        "[8] A. Niculescu-Mizil and R. Caruana, \"Predicting good probabilities with supervised learning,\" in Proc. 22nd International Conference on Machine Learning (ICML), Bonn, Germany, 2005, pp. 625–632.",
        "[9] D. W. Aha, \"Heart Disease Databases Documentation,\" Department of Information and Computer Science, University of California, Irvine, Tech. Rep., July 1988.",
        "[10] C. M. Bishop, Pattern Recognition and Machine Learning. New York, NY, USA: Springer, 2006."
    ]
    for r in refs:
        p = doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.05
        p.paragraph_format.space_after = Pt(2.5)
        p.paragraph_format.left_indent = Inches(0.2)
        p.paragraph_format.first_line_indent = Inches(-0.2)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        run = p.add_run(r)
        run.font.name = 'Times New Roman'
        run.font.size = Pt(8)

    doc.save('Final_CAD_Bayesian_Classification_Paper.docx')
    print("Successfully built Final_CAD_Bayesian_Classification_Paper.docx")

def convert_to_pdf_and_inject_high_res():
    import win32com.client
    import fitz

    docx_path = os.path.abspath('Final_CAD_Bayesian_Classification_Paper.docx')
    pdf_path = os.path.abspath('Final_CAD_Bayesian_Classification_Paper.pdf')

    word = win32com.client.Dispatch('Word.Application')
    word.Visible = False
    try:
        wdoc = word.Documents.Open(docx_path)
        # Export with wdExportOptimizeForPrint (0) for maximum base clarity
        wdoc.ExportAsFixedFormat(
            OutputFileName=pdf_path,
            ExportFormat=17, # wdExportFormatPDF
            OpenAfterExport=False,
            OptimizeFor=0,   # wdExportOptimizeForPrint
            BitmapMissingFonts=True,
            UseISO19005_1=False
        )
        wdoc.Close()
    finally:
        word.Quit()

    # Now inject full-resolution 300 DPI lossless images into the generated PDF
    pdoc = fitz.open(pdf_path)
    
    # Map of images in the paper in exact sequential order
    fig_files = [
        'figs/f2.png',
        'figs/f1.png',
        'figs/f3.png',
        'figs/f5.png',
        'figs/f6.png',
        'figs/f4.png',
        'figs/f7.png'
    ]
    
    img_list = []
    for pno, page in enumerate(pdoc):
        for img in page.get_images():
            img_list.append((pno, img[0]))

    if len(img_list) == len(fig_files):
        for (pno, xref), img_path in zip(img_list, fig_files):
            page = pdoc[pno]
            page.replace_image(xref=xref, filename=img_path)
        print(f"Successfully injected {len(fig_files)} uncompressed 300 DPI figures into PDF.")
    else:
        print(f"Warning: Found {len(img_list)} images in PDF, expected {len(fig_files)}.")

    pdoc.save(pdf_path, incremental=True, encryption=fitz.PDF_ENCRYPT_KEEP)
    pdoc.close()
    print("Successfully produced crystal-clear Final_CAD_Bayesian_Classification_Paper.pdf")

if __name__ == '__main__':
    build_paper()
    convert_to_pdf_and_inject_high_res()
